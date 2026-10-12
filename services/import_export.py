# -*- coding: utf-8 -*-
"""
services.import_export
ورود/خروج داده‌های صورتجلسات از/به CSV و Excel

نکات مهم هماهنگی ستون‌ها:
- نام ستون‌ها نرمال‌سازی می‌شود (فاصله/نیم‌فاصله، ی/ک عربی، اعداد فارسی، ...).
- برای هر فیلد چندین alias پشتیبانی می‌شود (مثلاً "تاریخ جلسه" و "تاریخ صورتجلسه").
- قبل از شروع واردات، وجود ستون‌های ضروری بررسی می‌شود و اگر نباشد خطا می‌دهد.
"""
from __future__ import annotations

import csv
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Tuple

from PyQt6.QtCore import QObject, pyqtSignal, QRunnable

from db.manager import DatabaseManager
from db.models import Meeting, DISCIPLINES
from utils.date_utils import to_gregorian

try:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ──────────────────────────────────────────────────────────────────
# Normalization helpers (برای هماهنگی نام ستون‌ها)
# ──────────────────────────────────────────────────────────────────

_DIGIT_TRANS = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)


def _as_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v)


def _norm(s: Any) -> str:
    """
    نرمال‌سازی کلید ستون برای مقایسه یکنواخت.
    - حذف نیم‌فاصله
    - یکسان‌سازی ي/ی و ك/ک
    - یکسان‌سازی اعداد فارسی/عربی به انگلیسی
    - حذف فاصله‌های اضافی و تبدیل جداکننده‌ها به underscore
    """
    s = _as_str(s)
    s = s.translate(_DIGIT_TRANS)
    s = s.replace("\u200c", "")  # ZWNJ
    s = s.replace("ي", "ی").replace("ك", "ک")
    s = s.replace("ة", "ه").replace("ۀ", "ه")
    s = s.strip()
    s = " ".join(s.split())

    # یکسان‌سازی جداکننده‌ها
    for ch in ("/", "\\", "-", "–", "—", ":", "؛", "،", "\t"):
        s = s.replace(ch, " ")

    parts = [p for p in s.split(" ") if p]
    return "_".join(parts).lower()


def _fix_numeric_string(s: str) -> str:
    """
    اگر مقدار از اکسل به شکل "123.0" آمده باشد، به "123" تبدیل می‌کند.
    """
    s = (s or "").strip()
    if not s:
        return ""
    if s.endswith(".0"):
        core = s[:-2]
        if core.replace("-", "").isdigit():
            return core
    return s


# ──────────────────────────────────────────────────────────────────
# Parsing helpers
# ──────────────────────────────────────────────────────────────────

def _parse_date(value: Any) -> Optional[date]:
    """
    تبدیل تاریخ (شمسی یا میلادی) به date.

    ترتیب بررسی:
    1) اگر date/datetime باشد
    2) تاریخ شمسی (YYYY-MM-DD با سال > 1300)
    3) فرمت‌های رایج میلادی
    4) isoformat (با/بدون زمان)
    """
    if value is None:
        return None

    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()

    s = _as_str(value).strip()
    if not s:
        return None

    # برخی سلول‌های اکسل ممکن است "YYYY-MM-DD 00:00:00" باشند
    s = s.replace("T", " ").strip()

    # تاریخ شمسی
    if s.count("-") >= 2 and len(s) >= 10 and s[:4].isdigit() and int(s[:4]) > 1300:
        sh = s[:10]
        try:
            g = to_gregorian(sh)
            if g:
                return datetime.strptime(g, "%Y-%m-%d").date()
        except Exception:
            pass

    # فرمت‌های میلادی
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue

    try:
        # ISO (ممکن است شامل زمان باشد)
        return datetime.fromisoformat(s.replace("Z", "")).date()
    except Exception:
        return None


def _clean_discipline(value: Any) -> Optional[str]:
    """
    پاک‌سازی/اعتبارسنجی دیسیپلین.
    اگر در لیست DISCIPLINES نبود، همان مقدار (trim شده) ذخیره می‌شود (انعطاف).
    """
    v = _as_str(value).strip()
    if not v:
        return None

    if v in DISCIPLINES:
        return v

    v_lower = v.lower()
    for d in DISCIPLINES:
        if d.lower() == v_lower:
            return d

    return v


# ──────────────────────────────────────────────────────────────────
# Export utilities (CSV and Excel)
# ──────────────────────────────────────────────────────────────────

def export_meetings_to_csv(path: str, headers: List[str], rows: List[List[Any]]) -> None:
    """صادرات لیست صورتجلسات به فایل CSV با BOM."""
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in rows:
            w.writerow(["" if v is None else _as_str(v) for v in r])


def export_meetings_to_excel(path: str, headers: List[str], rows: List[List[Any]]) -> None:
    """صادرات لیست صورتجلسات به فایل Excel."""
    if not HAS_OPENPYXL:
        raise ImportError("برای خروجی Excel باید openpyxl نصب باشد: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "Meetings"

    # هدر
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    # داده‌ها
    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value="" if value is None else _as_str(value))

    # عرض خودکار ستون‌ها (حداکثر ۵۰)
    for col_idx in range(1, len(headers) + 1):
        max_width = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
            for cell_value in row:
                if cell_value:
                    max_width = max(max_width, len(_as_str(cell_value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width + 2, 50)

    wb.save(path)


# ──────────────────────────────────────────────────────────────────
# Import worker for Excel/CSV
# ──────────────────────────────────────────────────────────────────

class ImportWorkerSignals(QObject):
    progress = pyqtSignal(int)   # 0..100
    finished = pyqtSignal(dict)  # {'added': int, 'skipped': List[str], 'cancelled': bool}
    error    = pyqtSignal(str)   # message


class ImportWorker(QRunnable):
    """
    Worker رشته‌ای برای واردات داده از فایل اکسل یا CSV.

    نکته: کلیدهای دیکشنری هر ردیف، نرمال‌شده هستند (با _norm).
    """

    # alias های هر فیلد (برای هماهنگی نام ستون‌ها)
    FIELD_ALIASES: Dict[str, Tuple[str, ...]] = {
        "meeting_number": ("meeting_number", "شماره صورتجلسه", "شماره جلسه"),
        "subject":        ("subject", "موضوع", "عنوان"),
        "meeting_date":   ("meeting_date", "تاریخ صورتجلسه", "تاریخ جلسه"),
        "due_date":       ("due_date", "مهلت انجام", "مهلت", "deadline"),

        "entry_date":     ("entry_date", "تاریخ ورود"),
        "employer":       ("employer", "کارفرما"),
        "consultant":     ("consultant", "مشاور"),
        "contractor":     ("contractor", "پیمانکار"),
        "drawing_number": ("drawing_number", "شماره نقشه"),
        "contract_number":("contract_number", "شماره قرارداد"),
        "contract_date":  ("contract_date", "تاریخ قرارداد"),
        "sent_to_contractor_date": ("sent_to_contractor_date", "تاریخ ارسال به پیمانکار", "تاریخ ارسال"),
        "status":         ("status", "وضعیت"),
        "priority":       ("priority", "اولویت"),
        "closing_date":   ("closing_date", "تاریخ بسته شدن", "تاریخ بسته‌شدن"),
        "file_link":      ("file_link", "لینک فایل"),
        "discipline":     ("discipline", "دیسیپلین", "تخصص", "رشته"),
    }

    REQUIRED_FIELDS: Tuple[str, ...] = ("meeting_number", "subject", "meeting_date", "due_date")

    def __init__(self, db: DatabaseManager, file_path: str) -> None:
        super().__init__()
        self.db = db
        self.file_path = file_path
        self.signals = ImportWorkerSignals()
        self._cancelled = False

    def cancel(self) -> None:
        """لغو واردات در اولین فرصت."""
        self._cancelled = True

    # ── اجرای اصلی ────────────────────────────────────────────────

    def run(self) -> None:
        try:
            rows, keys = self._read_rows_any_with_keys()
            self._validate_required_headers(keys)
        except Exception as e:
            self.signals.error.emit(str(e))
            return

        total = max(len(rows), 1)
        added = 0
        skipped: List[str] = []

        # برای جلوگیری از تکرار در همان فایل
        seen_numbers: set[str] = set()

        self.signals.progress.emit(0)

        with self.db.session() as sess:
            for i, row in enumerate(rows, start=1):
                if self._cancelled:
                    break

                self.signals.progress.emit(int((i / total) * 100))

                excel_row_no = i + 1  # چون ردیف 1 هدر است و داده از ردیف 2 شروع می‌شود

                try:
                    meeting_number = _fix_numeric_string(
                        (self._pick(row, *self.FIELD_ALIASES["meeting_number"]) or "").strip()
                    )
                    subject = (self._pick(row, *self.FIELD_ALIASES["subject"]) or "").strip()

                    meeting_date_str = self._pick(row, *self.FIELD_ALIASES["meeting_date"])
                    due_date_str     = self._pick(row, *self.FIELD_ALIASES["due_date"])

                    employer   = (self._pick(row, *self.FIELD_ALIASES["employer"]) or "").strip()
                    consultant = (self._pick(row, *self.FIELD_ALIASES["consultant"]) or "").strip()
                    contractor = (self._pick(row, *self.FIELD_ALIASES["contractor"]) or "").strip()

                    drawing_number  = (self._pick(row, *self.FIELD_ALIASES["drawing_number"]) or "").strip()
                    contract_number = (self._pick(row, *self.FIELD_ALIASES["contract_number"]) or "").strip()

                    contract_date_str = self._pick(row, *self.FIELD_ALIASES["contract_date"])
                    sent_date_str     = self._pick(row, *self.FIELD_ALIASES["sent_to_contractor_date"])
                    closing_date_str  = self._pick(row, *self.FIELD_ALIASES["closing_date"])

                    status = (self._pick(row, *self.FIELD_ALIASES["status"]) or "در دست بررسی").strip()
                    priority = (self._pick(row, *self.FIELD_ALIASES["priority"]) or "متوسط").strip()

                    file_link = (self._pick(row, *self.FIELD_ALIASES["file_link"]) or "").strip()

                    discipline_raw = self._pick(row, *self.FIELD_ALIASES["discipline"])
                    discipline = _clean_discipline(discipline_raw)

                    # ── اعتبارسنجی فیلدهای اجباری ─────────────────
                    if not meeting_number:
                        skipped.append(f"Row {excel_row_no}: مقدار «شماره صورتجلسه» خالی است")
                        continue
                    if not subject:
                        skipped.append(f"Row {excel_row_no}: مقدار «موضوع» خالی است")
                        continue
                    if not meeting_date_str:
                        skipped.append(f"Row {excel_row_no}: مقدار «تاریخ صورتجلسه» خالی است")
                        continue
                    if not due_date_str:
                        skipped.append(f"Row {excel_row_no}: مقدار «مهلت انجام» خالی است")
                        continue

                    # ── جلوگیری از تکرار داخل فایل ────────────────
                    if meeting_number in seen_numbers:
                        skipped.append(f"Row {excel_row_no}: شماره صورتجلسه تکراری در فایل: '{meeting_number}'")
                        continue
                    seen_numbers.add(meeting_number)

                    # ── تبدیل تاریخ‌ها ────────────────────────────
                    meeting_date  = _parse_date(meeting_date_str)
                    due_date      = _parse_date(due_date_str)
                    contract_date = _parse_date(contract_date_str) if contract_date_str else None
                    sent_date     = _parse_date(sent_date_str)     if sent_date_str     else None
                    closing_date  = _parse_date(closing_date_str)  if closing_date_str  else None

                    if not meeting_date:
                        skipped.append(f"Row {excel_row_no}: تاریخ صورتجلسه نامعتبر: '{_as_str(meeting_date_str)}'")
                        continue
                    if not due_date:
                        skipped.append(f"Row {excel_row_no}: مهلت انجام نامعتبر: '{_as_str(due_date_str)}'")
                        continue

                    # ── بررسی تکراری بودن در دیتابیس ──────────────
                    exists = (
                        sess.query(Meeting.id)
                        .filter(Meeting.meeting_number == meeting_number)
                        .first()
                    )
                    if exists:
                        skipped.append(f"Row {excel_row_no}: شماره صورتجلسه تکراری در دیتابیس: '{meeting_number}'")
                        continue

                    # شمارش پیوست
                    attachment_count = 1 if file_link else 0

                    # entry_date (در صورت نبود، امروز)
                    entry_date_str = self._pick(row, *self.FIELD_ALIASES["entry_date"])
                    entry_date = _parse_date(entry_date_str) if entry_date_str else None
                    entry_date = entry_date or date.today()

                    # ── ایجاد رکورد ───────────────────────────────
                    m = Meeting(
                        meeting_number=meeting_number,
                        subject=subject,
                        meeting_date=meeting_date,
                        due_date=due_date,
                        employer=employer or None,
                        consultant=consultant or None,
                        contractor=contractor or None,
                        drawing_number=drawing_number or None,
                        contract_number=contract_number or None,
                        contract_date=contract_date,
                        sent_to_contractor_date=sent_date,
                        status=status,
                        priority=priority,
                        closing_date=closing_date,
                        file_link=file_link or None,
                        attachment_count=attachment_count,
                        entry_date=entry_date,
                        discipline=discipline,
                    )

                    sess.add(m)
                    sess.commit()
                    added += 1

                except Exception as exc:
                    sess.rollback()
                    skipped.append(f"Row {excel_row_no}: {exc}")

        self.signals.progress.emit(100)
        self.signals.finished.emit({
            "added":     added,
            "skipped":   skipped,
            "cancelled": self._cancelled,
        })

    # ── helpers ───────────────────────────────────────────────────

    def _validate_required_headers(self, keys: set[str]) -> None:
        """
        بررسی می‌کند ستون‌های ضروری (با aliasها) در فایل وجود داشته باشند.
        """
        missing: List[str] = []
        for field in self.REQUIRED_FIELDS:
            aliases = self.FIELD_ALIASES.get(field, (field,))
            if not any(_norm(a) in keys for a in aliases):
                # اسم نمایشی را همان اولین alias فارسی/اصلی می‌گذاریم
                missing.append(aliases[1] if len(aliases) > 1 else aliases[0])

        if missing:
            raise ValueError(
                "ستون‌های ضروری در فایل پیدا نشد:\n- "
                + "\n- ".join(missing)
                + "\n\n(نام ستون‌ها را دقیقاً مطابق قالب/خروجی برنامه تنظیم کنید.)"
            )

    @staticmethod
    def _pick(row: Dict[str, Any], *keys: str) -> str:
        """
        دریافت اولین مقدار معتبر از دیکشنری ردیف با کلیدهای مختلف.
        کلیدها ابتدا نرمال‌سازی می‌شوند.
        """
        for k in keys:
            nk = _norm(k)
            if nk in row and row.get(nk) is not None:
                return _as_str(row[nk])
        return ""

    def _read_rows_any_with_keys(self) -> Tuple[List[Dict[str, Any]], set[str]]:
        """
        خواندن ردیف‌ها از فایل Excel یا CSV + مجموعه کلیدهای نرمال‌شده.
        """
        path_lower = self.file_path.lower()
        if path_lower.endswith((".xlsx",)):
            rows = self._read_excel()
        elif path_lower.endswith((".xls",)):
            # .xls معمولاً نیاز به xlrd دارد؛ برای جلوگیری از رفتارهای نامطمئن صریحاً خطا می‌دهیم.
            raise ValueError(
                "فرمت .xls پشتیبانی نمی‌شود. لطفاً فایل را به .xlsx تبدیل کنید."
            )
        else:
            rows = self._read_csv()

        keys: set[str] = set()
        for r in rows:
            keys.update(r.keys())
        return rows, keys

    def _read_excel(self) -> List[Dict[str, Any]]:
        """
        خواندن فایل اکسل با pandas (پیشنهادی برای سازگاری بهتر).
        """
        try:
            import pandas as pd
        except ImportError as exc:
            raise ValueError(
                "برای واردات Excel باید pandas نصب باشد.\n"
                "نصب: pip install pandas openpyxl"
            ) from exc

        try:
            df = pd.read_excel(self.file_path, dtype=object, engine="openpyxl")
        except Exception as exc:
            raise ValueError(f"خواندن فایل اکسل ناموفق بود: {exc}") from exc

        df = df.where(pd.notna(df), None)

        # rename columns to normalized keys
        df.rename(columns={c: _norm(c) for c in df.columns}, inplace=True)

        out: List[Dict[str, Any]] = []
        for _, r in df.iterrows():
            d = r.to_dict()
            out.append({str(k): ("" if v is None else _as_str(v)) for k, v in d.items()})

        return out

    def _read_csv(self) -> List[Dict[str, Any]]:
        """
        خواندن فایل CSV با نگاشت کلید نرمال‌شده.
        """
        try:
            with open(self.file_path, "r", encoding="utf-8-sig", newline="") as f:
                reader = csv.DictReader(f)
                if not reader.fieldnames:
                    return []

                norm_map = {_norm(h): h for h in reader.fieldnames}
                out: List[Dict[str, Any]] = []

                for row in reader:
                    normalized: Dict[str, Any] = {}
                    for nk, orig in norm_map.items():
                        normalized[nk] = row.get(orig)
                    out.append(normalized)

            return out
        except Exception as exc:
            raise ValueError(f"خواندن فایل CSV ناموفق بود: {exc}") from exc